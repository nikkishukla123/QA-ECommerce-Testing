from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path


OUT_DIR = Path(r"C:\New project\QA_Ecommerce_Testing_Project")
DOCX_PATH = OUT_DIR / "E-Commerce_QA_Project_Full_Report_and_Interview_Guide.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_twips):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = OxmlElement("w:tcW")
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")
    tc_pr.append(tc_w)


def set_table_borders(table, color="D1D5DB", size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def style_doc(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    title = doc.styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(31, 58, 95)

    for style_name, size, color in [
        ("Heading 1", 16, RGBColor(46, 116, 181)),
        ("Heading 2", 13, RGBColor(46, 116, 181)),
        ("Heading 3", 12, RGBColor(31, 77, 120)),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)


def add_title(doc):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("E-Commerce Website Testing Project")
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p2.add_run("Full Project Report + Interview Preparation Guide")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(55, 65, 81)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.add_run("Modules: Registration, Login, Search, Product Details, Add to Cart, Checkout, Payment, Order Confirmation")
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.add_run("Testing Types: Functional, Regression, Smoke, Negative, UI, Boundary Value Testing")


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, "9CA3AF", "6")
    cell = table.cell(0, 0)
    set_cell_shading(cell, "E0F2FE")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    r = p.add_run(title + ": ")
    r.bold = True
    r.font.color.rgb = RGBColor(31, 58, 95)
    p.add_run(body)


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = text
        set_cell_shading(hdr[i], "1F3A5F")
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in hdr[i].paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
        if widths:
            set_cell_width(hdr[i], widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = str(text)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                set_cell_width(cells[i], widths[i])
    return table


def add_qa_section(doc, title, questions):
    doc.add_heading(title, level=1)
    for idx, (q, a) in enumerate(questions, 1):
        p = doc.add_paragraph()
        r = p.add_run(f"Q{idx}. {q}")
        r.bold = True
        r.font.color.rgb = RGBColor(31, 77, 120)
        doc.add_paragraph(a)


basic_questions = [
    ("Manual testing kya hota hai?", "Manual testing mein tester application ko manually use karke verify karta hai ki actual result expected result ke according hai ya nahi. Is project mein maine user journeys jaise registration, login, search, cart aur checkout manually execute kiye."),
    ("Functional testing kya hoti hai?", "Functional testing mein hum verify karte hain ki application requirement ke according kaam kar rahi hai. Example: valid login credentials dene par user dashboard/home page par jana chahiye."),
    ("Regression testing kya hoti hai?", "Regression testing mein bug fix ya new build ke baad existing functionality retest karte hain. Is project mein cart, login, search aur checkout flows ko regression mein cover kiya."),
    ("Smoke testing kya hoti hai?", "Smoke testing basic build verification hoti hai. Isme check kiya ki major flows open ho rahe hain ya nahi, jaise login page load, product search, add to cart, checkout page open."),
    ("Test case kya hota hai?", "Test case ek documented step-by-step condition hoti hai jisme test data, steps, expected result, actual result aur status mention hota hai."),
    ("Test scenario aur test case mein difference?", "Test scenario high-level condition hoti hai, jaise 'Verify login functionality'. Test case detailed steps hota hai, jaise valid email/password enter karke login button click karna."),
    ("Severity aur priority mein difference?", "Severity bug ka technical/business impact batati hai. Priority batati hai bug ko kitni jaldi fix karna hai. Example: checkout empty address allow karta hai, severity critical aur priority critical."),
    ("Bug life cycle kya hota hai?", "Typical bug life cycle: New > Assigned > Open > Fixed > Retest > Verified > Closed. Agar issue fix nahi hua to Reopen hota hai."),
    ("Positive testing kya hai?", "Valid data se functionality verify karna positive testing hai. Example: valid email aur password se login successful hona."),
    ("Negative testing kya hai?", "Invalid input dekar validation verify karna negative testing hai. Example: invalid email format, blank password, alphabetic CVV."),
    ("Boundary value testing kya hai?", "Input field ke minimum/maximum boundaries test karna. Example: PIN code length, mobile number digits, password min length."),
    ("UI testing kya cover karta hai?", "UI testing labels, alignment, buttons, error messages, images, responsiveness aur visibility verify karta hai."),
    ("Retesting aur regression testing mein difference?", "Retesting sirf fixed bug ko dobara verify karta hai. Regression testing related aur existing functionality ko verify karta hai ki fix se kuch break nahi hua."),
    ("Test data kya use kiya?", "Valid/invalid email, valid/invalid password, existing user, invalid mobile, product keywords, valid/invalid PIN, valid/invalid coupon, valid/invalid card details."),
    ("Defect report mein kya fields hote hain?", "Bug ID, title, module, severity, priority, environment, steps to reproduce, expected result, actual result, attachment, status, assignee."),
    ("JIRA ka use kaise kiya?", "Mock JIRA tickets create kiye jisme issue type bug, summary, steps, expected/actual result, severity, priority aur status maintain kiya."),
    ("Requirement Traceability Matrix kya hota hai?", "RTM requirements ko test cases se map karta hai, jisse coverage prove hoti hai. Example: REQ-005 cart requirement TC-029 to TC-039 se covered hai."),
    ("Exit criteria kya tha?", "All planned test cases executed, critical/high bugs logged, regression completed, and test summary prepared."),
    ("Entry criteria kya tha?", "Application accessible, build stable, core modules available, test data ready, requirements understood."),
    ("Pass/fail kaise decide kiya?", "Expected result aur actual result compare karke. Match hua to Pass, mismatch hua to Fail and bug logged."),
]


project_questions = [
    ("Apne project ke baare mein batao.", "Maine Amazon/Flipkart jaisi e-commerce application test ki. Scope mein registration, login, search, product details, cart, checkout, payment validation aur order confirmation tha. Maine 55 test cases create aur execute kiye, 10 bugs identify kiye, mock JIRA defects banaye, RTM prepare kiya aur regression testing perform ki."),
    ("Aapka role kya tha?", "Mera role manual QA tester ka tha. Maine requirements samjhe, test scenarios banaye, test cases design kiye, execution status maintain kiya, bugs report kiye, bug fixes retest kiye aur regression testing perform ki."),
    ("Aapne test cases kaise design kiye?", "Maine modules ko break kiya: Registration, Login, Search, Product Details, Cart, Checkout, Payment, Order Confirmation. Har module ke liye positive, negative, boundary, UI aur regression test cases likhe."),
    ("Kaunse important bugs mile?", "Important bugs: cart quantity update par subtotal refresh nahi ho raha tha, checkout empty address accept kar raha tha, removed cart item refresh ke baad wapas aa raha tha, invalid PIN code accepted tha, CVV alphabetic characters accept kar raha tha."),
    ("Sabse critical bug kaunsa tha?", "Checkout mein empty address ke saath Place Order button enabled tha. Ye critical tha kyunki user mandatory shipping details ke bina order attempt kar sakta tha."),
    ("Bug reproduce kaise karte the?", "Main exact steps note karta tha: login, product add, checkout open, address fields blank, Place Order observe. Same steps repeat karke bug confirm karta tha."),
    ("Aapne regression suite kaise select kiya?", "Regression ke liye critical business flows select kiye: login, search product, add to cart, update cart, checkout, logout. Bug impacted areas ko bhi include kiya."),
    ("Search module mein kya test kiya?", "Valid keyword, invalid keyword, blank search, suggestions, product count, price filter, sorting, clear filter behavior test kiya."),
    ("Cart module mein kya test kiya?", "Add to cart, duplicate product, cart count, quantity update, remove item, refresh persistence, multiple products, out-of-stock product, coupon apply, proceed to checkout test kiya."),
    ("Checkout module mein kya test kiya?", "Checkout page load, valid address, blank address validation, city validation, PIN validation, saved address select/edit, order summary total verify kiya."),
    ("Payment gateway real tha?", "Project scope mein real payment gateway processing out of scope tha. Payment validation sandbox/mock level par test ki, jaise card number, expiry, CVV validations."),
    ("Aapne expected result kaise likha?", "Requirement aur standard e-commerce behavior ke basis par. Example: mandatory address blank hone par validation message aana chahiye aur Place Order disabled hona chahiye."),
    ("Aapne actual result kaise maintain kiya?", "Execution ke time actual application behavior note kiya. Agar expected se match karta tha to Pass, mismatch hota tha to Fail aur bug ID link kiya."),
    ("Aapne priority kaise assign ki?", "Business urgency ke basis par. Checkout blocking issue critical priority, UI image issue high/medium depending impact, minor validation inconsistency medium/low."),
    ("Project ka final result kya tha?", "55 test cases execute hue, 45 pass aur 10 fail/bugs logged. Critical/high issues documented hue and regression impact areas identify kiye gaye."),
]


advanced_questions = [
    ("Agar requirement clear na ho to kya karoge?", "Main requirement clarification raise karunga, acceptance criteria ask karunga, and assumption document karunga. Jab tak clarification nahi milti, main standard behavior ke basis par test scenario draft kar sakta hoon but final status requirement confirmation ke baad hi close karunga."),
    ("Cart subtotal bug frontend hai ya backend, kaise identify karoge?", "Network tab se API response check karunga. Agar API correct subtotal return kar rahi hai but UI update nahi ho raha, frontend issue. Agar API hi old total bhej rahi hai, backend/calculation issue."),
    ("Removed item refresh ke baad wapas aa raha hai, root cause kya ho sakta hai?", "Possibilities: remove API fail ho rahi hai, database/cart state update nahi ho raha, frontend local state update ho raha but server state stale hai, cache issue hai, ya session cart sync issue hai."),
    ("Checkout empty address allow kar raha hai, kya risk hai?", "Order processing failure, wrong/incomplete shipping, customer dissatisfaction, support tickets, revenue leakage, and compliance/operational risk."),
    ("Aap browser devtools kaise use karoge?", "Console errors, network API status, request payload, response body, local storage/session storage, cookies, and UI errors inspect karunga."),
    ("Agar bug developer ke system mein reproduce nahi hota to?", "Environment details, browser version, user account, test data, screenshots/video, exact steps share karunga. Fresh session/incognito and different browser se reproduce proof provide karunga."),
    ("Flaky bug kaise handle karte ho?", "Multiple attempts, pattern observation, logs, timestamps, network conditions, exact data capture, and frequency mention karta hoon. Bug title mein intermittent behavior clearly mention karunga."),
    ("Production mein critical checkout bug mil jaye to kya karoge?", "Immediately severity critical mark karunga, lead/manager ko notify karunga, reproduction proof attach karunga, impact mention karunga, hotfix validation plan prepare karunga."),
    ("Regression testing kitna broad hona chahiye?", "Change impact analysis ke basis par. Agar cart fix hai to cart, checkout, order summary, coupon, payment flow regression cover karunga. Full regression tab hota hai jab shared components ya major build change ho."),
    ("API testing include karoge?", "Agar access mile to haan. Login API, search API, cart add/update/remove API, checkout address API validate karunga for status code, payload, response, error handling."),
    ("SQL/database verification kab useful hai?", "Order creation, cart persistence, user registration, address save, coupon usage jaise backend state validations ke liye useful hai, agar tester ko DB read access allowed ho."),
    ("Security angle se kya basic checks karoge?", "Password masking, invalid login error, session expiry, logout session clear, direct checkout URL access, input validation, sensitive data not visible in URL/logs."),
    ("Performance testing is project mein ki?", "Formal performance testing out of scope tha, but manual observation mein page load, search response, cart update delay note kar sakte hain. Dedicated performance testing ke liye JMeter/Lighthouse type tools use honge."),
    ("Mobile responsiveness kaise test karoge?", "Chrome responsive mode and real device if available. Header, search bar, product cards, cart buttons, checkout form fields, and payment section layout verify karunga."),
    ("Agar 50 test cases ke liye time kam ho to priority kaise set karoge?", "Risk-based testing karunga. First critical revenue flows: login, search, add to cart, checkout, payment validation, order confirmation. Then negative and UI cases."),
    ("Defect leakage kaise reduce karoge?", "Clear requirements, peer review of test cases, risk-based coverage, exploratory testing, regression suite, bug trend analysis, and production-like test data se."),
    ("Test coverage kaise prove karoge?", "RTM se requirements to test case mapping dikhakar. Dashboard mein total test cases, pass/fail, bugs by severity and module summary maintain karke."),
    ("What is impact analysis?", "Code change or bug fix se affected modules identify karna. Example: cart quantity fix can impact subtotal, coupon discount, checkout order summary, and payment amount."),
    ("Exploratory testing kya ki?", "Structured test cases ke alawa user ki tarah application explore ki: repeated add/remove, refresh, invalid inputs, browser back button, blank search, coupon retry."),
    ("Aap automation ke liye kaunse cases choose karoge?", "Stable, repetitive, high-value regression cases: login, search, add to cart, cart update, checkout validation, logout. Frequently changing UI ya one-time cases manual rakhenge."),
]


tricky_questions = [
    ("Agar developer bole ye bug nahi requirement hai, kya karoge?", "Main requirement/acceptance criteria refer karunga. Agar documented nahi hai to product owner/BA se clarification lunga. Discussion professional rakhta hoon and evidence ke basis par decision accept karta hoon."),
    ("Aapne 55 test cases kyun banaye, 20 kyun nahi?", "Modules multiple the aur har module mein positive, negative, boundary, UI and regression scenarios required the. 55 test cases se core customer journey and validations better cover hue."),
    ("Kya aapne real Amazon/Flipkart test kiya?", "Nahi, ye Amazon/Flipkart jaisi demo/e-commerce application par manual testing project tha. Maine real e-commerce behavior ko reference karke test scenarios design kiye."),
    ("Agar interviewer bole project fake lag raha hai?", "Main confidently artifacts dikhaunga: test plan, 55 test cases, bug report, mock JIRA tickets, RTM, execution dashboard. Main exact bugs aur reproduce steps explain kar sakta hoon."),
    ("Aapko bug ka root cause pata tha?", "As QA, main root cause confirm nahi karta jab tak logs/code access na ho. Main probable area mention kar sakta hoon, but bug report mein observed behavior, steps, expected/actual and impact clearly deta hoon."),
    ("Kya har fail test case bug hota hai?", "Mostly fail test case defect indicate karta hai, but kabhi requirement misunderstanding, environment issue, test data issue ya expected result wrong ho sakta hai. Isliye triage important hai."),
    ("Agar same bug multiple modules mein mile to?", "Main duplicate avoid karunga. Ek parent bug create karke affected modules mention karunga, ya separate bugs create karunga if root/impact alag hai."),
    ("Kya low severity bug high priority ho sakta hai?", "Haan. Example: spelling mistake on home page before important release low severity but high priority ho sakti hai due to brand impact."),
    ("Kya high severity bug low priority ho sakta hai?", "Haan. Example: critical crash only in rare unsupported browser, severity high but priority low ho sakti hai if business impact low hai."),
    ("Interview mein project ka best answer kaise start karoge?", "I tested an e-commerce web application similar to Amazon/Flipkart. I prepared a test plan, wrote and executed 55 manual test cases, logged 10 defects in mock JIRA format, and performed functional and regression testing across login, search, cart, checkout, and payment validation."),
]


def build():
    doc = Document()
    style_doc(doc)
    add_title(doc)
    add_callout(
        doc,
        "One-line interview pitch",
        "I tested an e-commerce website similar to Amazon/Flipkart, created and executed 55 manual test cases, logged 10 defects using mock JIRA, and performed functional plus regression testing for critical user journeys.",
    )

    doc.add_heading("1. Project Overview", level=1)
    doc.add_paragraph(
        "This project is a complete manual testing portfolio project for an e-commerce website. "
        "The application behavior is similar to Amazon/Flipkart, where users can register, login, search products, view product details, add products to cart, proceed to checkout, validate payment fields, and receive order confirmation."
    )
    add_bullets(
        doc,
        [
            "Project type: Manual QA / Functional Testing project",
            "Application domain: E-commerce",
            "Total test cases created: 55",
            "Bugs documented: 10",
            "Tools used: Excel, Mock JIRA, Browser DevTools",
            "Main testing types: Functional, Regression, Smoke, Negative, UI, Boundary Value Testing",
        ],
    )

    doc.add_heading("2. Objective", level=1)
    add_bullets(
        doc,
        [
            "Verify that core e-commerce user flows work according to expected behavior.",
            "Validate positive and negative scenarios for login, registration, search, cart, checkout, and payment fields.",
            "Identify defects that can affect customer experience, revenue flow, or order processing.",
            "Prepare professional QA artifacts that can be shown in a resume or interview.",
        ],
    )

    doc.add_heading("3. Scope of Testing", level=1)
    add_table(
        doc,
        ["In Scope", "Out of Scope"],
        [
            ["Registration and login validation", "Real payment settlement"],
            ["Search, filters, sorting, product listing", "Backend performance/load testing"],
            ["Product details and delivery PIN validation", "Security penetration testing"],
            ["Cart add/update/remove and coupons", "Real warehouse/inventory integration"],
            ["Checkout address and payment field validation", "Production deployment validation"],
        ],
        [4680, 4680],
    )

    doc.add_heading("4. Test Environment", level=1)
    add_table(
        doc,
        ["Area", "Details"],
        [
            ["Browser", "Chrome and Edge"],
            ["Device", "Desktop and mobile responsive view"],
            ["Operating System", "Windows"],
            ["Test data", "Valid/invalid email, password, product keywords, address, PIN, coupon code, card data"],
            ["Defect tool", "Mock JIRA"],
            ["Test case tool", "Excel workbook"],
        ],
        [2200, 7160],
    )

    doc.add_heading("5. My Role and Responsibilities", level=1)
    add_bullets(
        doc,
        [
            "Understood the e-commerce workflow and prepared module-wise test scenarios.",
            "Created 55 manual test cases with steps, expected results, actual results, status, priority, and type.",
            "Executed test cases and compared expected vs actual behavior.",
            "Logged defects with severity, priority, reproduction steps, expected result, actual result, and impact.",
            "Prepared RTM to show requirement coverage.",
            "Performed regression testing for critical flows after defect identification.",
            "Prepared final test summary and resume-ready project explanation.",
        ],
    )

    doc.add_heading("6. STLC Followed in This Project", level=1)
    add_numbered(
        doc,
        [
            "Requirement analysis: I understood core e-commerce requirements and user journeys.",
            "Test planning: I defined scope, environment, entry criteria, exit criteria, risks, and deliverables.",
            "Test case design: I created positive, negative, boundary, UI, smoke, and regression test cases.",
            "Test environment setup: Browser, test data, accounts, product keywords, and mock JIRA format were prepared.",
            "Test execution: I executed all 55 test cases and marked status as Pass/Fail.",
            "Defect reporting: I logged 10 bugs with severity, priority, and detailed reproduction steps.",
            "Retesting and regression: Failed scenarios and impacted flows were retested/regression-tested conceptually.",
            "Test closure: I prepared dashboard, bug report, RTM, and interview explanation.",
        ],
    )

    doc.add_heading("7. Module-wise Testing Summary", level=1)
    add_table(
        doc,
        ["Module", "What I Tested"],
        [
            ["Registration", "Valid registration, duplicate email, invalid email, blank mandatory fields, password mismatch, mobile number validation."],
            ["Login", "Valid login, invalid password, unregistered email, blank fields, forgot password link, logout, email space handling."],
            ["Search", "Valid keyword, invalid keyword, blank search, suggestions, result count, price filter, sorting, clear filter."],
            ["Product Details", "Page opening, image load, price visibility, rating/reviews, delivery PIN check, invalid PIN validation."],
            ["Cart", "Add item, duplicate item, cart count, quantity update, remove item, refresh persistence, coupons, out-of-stock item, checkout navigation."],
            ["Checkout", "Page load, valid address, blank address, city validation, PIN validation, saved address, edit address, order summary total."],
            ["Payment", "Valid card, invalid card, CVV digits-only validation, expired card validation."],
            ["Order Confirmation", "Confirmation page and order ID display after successful checkout."],
            ["Regression", "End-to-end login-search-cart-checkout flow and impacted areas after bug fixes."],
        ],
        [1900, 7460],
    )

    doc.add_heading("8. Test Design Techniques Used", level=1)
    add_table(
        doc,
        ["Technique", "How I Used It"],
        [
            ["Positive testing", "Used valid credentials, valid address, valid product keyword, valid coupon/card test data."],
            ["Negative testing", "Used invalid email, wrong password, blank address, invalid PIN, alphabetic CVV."],
            ["Boundary value testing", "Checked input length validations such as PIN code and mobile number."],
            ["UI testing", "Verified product images, field labels, buttons, messages, visibility, and alignment."],
            ["Smoke testing", "Verified major pages and critical actions open/work in the build."],
            ["Regression testing", "Rechecked login, search, cart, checkout, and logout after issues were identified."],
            ["Exploratory testing", "Tried refresh, repeated add/remove, invalid coupons, blank fields, and alternate navigation."],
        ],
        [2200, 7160],
    )

    doc.add_heading("9. Key Defects Found", level=1)
    add_table(
        doc,
        ["Bug ID", "Module", "Issue", "Severity", "Priority"],
        [
            ["BUG-001", "Registration", "Password mismatch error is not displayed", "Medium", "High"],
            ["BUG-003", "Search", "Search result count does not match displayed products", "Medium", "Medium"],
            ["BUG-005", "Cart", "Quantity update does not refresh total price immediately", "High", "High"],
            ["BUG-006", "Cart", "Removed item reappears after page refresh", "High", "Critical"],
            ["BUG-008", "Checkout", "Place Order button remains enabled with empty address", "Critical", "Critical"],
            ["BUG-009", "Payment", "Card CVV accepts alphabetic characters", "High", "High"],
            ["BUG-010", "Order Confirmation", "Confirmation page does not display order ID", "Medium", "High"],
        ],
        [1200, 1700, 3860, 1300, 1300],
    )

    doc.add_heading("10. Detailed Bug Explanation for Interview", level=1)
    add_callout(
        doc,
        "Critical bug example",
        "Checkout allowed the Place Order button even when mandatory address fields were empty. I reproduced it by logging in, adding a product, opening checkout, keeping address blank, and observing that the button was still enabled. Expected behavior was that validation should appear and the button should remain disabled.",
    )
    doc.add_paragraph(
        "This bug is critical because checkout is a revenue-critical flow. If users can attempt orders without shipping details, order processing can fail, support tickets can increase, and customer experience can be badly impacted."
    )

    doc.add_heading("11. Defect Reporting Format Used", level=1)
    add_table(
        doc,
        ["Field", "Example"],
        [
            ["Bug ID", "BUG-008"],
            ["Title", "Place Order button remains enabled with empty address"],
            ["Module", "Checkout"],
            ["Severity/Priority", "Critical/Critical"],
            ["Steps", "Login > Add product > Go to checkout > Keep address blank > Observe Place Order"],
            ["Expected", "Button disabled and validation displayed"],
            ["Actual", "Button enabled even with blank mandatory fields"],
            ["Impact", "Order can be attempted without valid shipping details"],
        ],
        [2200, 7160],
    )

    doc.add_heading("12. Regression Testing Explanation", level=1)
    doc.add_paragraph(
        "Regression testing was planned around critical e-commerce flows and impacted areas. For example, if cart quantity update is fixed, I would not only retest that one bug but also verify cart subtotal, coupon discount, checkout order summary, payment amount, and order confirmation amount."
    )
    add_bullets(
        doc,
        [
            "Login should still work after bug fixes.",
            "Search should still show correct product results.",
            "Add to cart should still update cart count.",
            "Cart quantity and remove item should work after refresh.",
            "Checkout should validate address and show correct order summary.",
            "Payment validation should still block invalid inputs.",
            "Logout/session behavior should remain correct.",
        ],
    )

    doc.add_heading("13. Metrics and Final Summary", level=1)
    add_table(
        doc,
        ["Metric", "Value"],
        [
            ["Total test cases", "55"],
            ["Passed test cases", "45"],
            ["Failed test cases / defects", "10"],
            ["Critical defects", "1"],
            ["High impact defects", "5+"],
            ["Main risk areas", "Cart persistence, checkout validation, payment field validation"],
            ["Deliverables", "Test plan, test cases, bug report, mock JIRA tickets, RTM, dashboard, interview guide"],
        ],
        [2600, 6760],
    )

    doc.add_heading("14. Resume and Interview Project Description", level=1)
    doc.add_heading("Short Resume Description", level=2)
    doc.add_paragraph(
        "Tested an e-commerce web application similar to Amazon/Flipkart covering registration, login, product search, add to cart, checkout, and order confirmation workflows. Created and executed 55 manual test cases, identified defects using mock JIRA, and performed functional and regression testing."
    )
    doc.add_heading("Resume Bullet Points", level=2)
    add_bullets(
        doc,
        [
            "Created and executed 55 manual test cases for registration, login, search, add to cart, checkout, payment validation, and order confirmation modules.",
            "Identified and documented defects with severity, priority, reproduction steps, expected results, actual results, and business impact.",
            "Performed functional, regression, smoke, negative, UI, and boundary value testing for critical e-commerce workflows.",
            "Prepared test plan, bug report, mock JIRA tickets, requirement traceability matrix, and test execution summary.",
        ],
    )
    doc.add_heading("30-second Interview Answer", level=2)
    doc.add_paragraph(
        "In my project, I tested an e-commerce website similar to Amazon/Flipkart. I covered modules like registration, login, product search, product details, cart, checkout, payment validation, and order confirmation. I created and executed 55 manual test cases, found 10 defects, documented them in mock JIRA format, and performed functional and regression testing."
    )
    doc.add_heading("2-minute Interview Answer", level=2)
    doc.add_paragraph(
        "This was a manual testing project for an e-commerce application. First, I understood the main user journey: a user registers or logs in, searches for products, checks product details, adds items to cart, applies coupons, proceeds to checkout, enters address/payment details, and receives order confirmation. Based on this flow, I prepared a test plan and wrote 55 test cases. I covered positive, negative, boundary, UI, smoke, and regression scenarios. During execution, I found issues like cart subtotal not updating after quantity change, removed cart item reappearing after refresh, checkout allowing empty address, invalid PIN accepted, and CVV accepting alphabets. I documented these bugs with severity, priority, steps to reproduce, expected result, actual result, and impact. I also prepared RTM and a final dashboard to show coverage and execution status."
    )

    add_qa_section(doc, "15. Basic Manual Testing Interview Questions", basic_questions)
    add_qa_section(doc, "16. Project-specific Interview Questions", project_questions)
    add_qa_section(doc, "17. Advanced and Scenario-based Interview Questions", advanced_questions)
    add_qa_section(doc, "18. Tricky Interview Questions", tricky_questions)

    doc.add_heading("19. Strong Closing Lines for Interview", level=1)
    add_bullets(
        doc,
        [
            "I can explain this project module-wise with test cases and bug examples.",
            "My strongest area in this project is functional testing and defect documentation.",
            "I understand the difference between retesting and regression through real cart and checkout examples.",
            "I can explain severity, priority, bug life cycle, RTM, and JIRA defect reporting using this project.",
            "If given application access, I can create test scenarios, execute cases, log defects, and prepare a test summary.",
        ],
    )

    doc.add_heading("20. What to Show If Interviewer Asks for Proof", level=1)
    add_bullets(
        doc,
        [
            "Excel workbook: E-Commerce_Website_Testing_Project.xlsx",
            "Test plan: Test_Plan.md",
            "Bug report: Bug_Report.md",
            "Mock JIRA tickets: Mock_JIRA_Tickets.md",
            "RTM: Requirement_Traceability_Matrix.md",
            "Resume content: Resume_Content.md",
            "This full report and interview guide document.",
        ],
    )

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build()
