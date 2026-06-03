from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path


OUT = Path(r"C:\New project\QA_Ecommerce_Testing_Project\Nikki_Shukla_QA_Resume_One_Page.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9E2F3"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_margins(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)


def set_base_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(9.4)
    normal.paragraph_format.space_after = Pt(1.5)
    normal.paragraph_format.line_spacing = 1.03

    for name, size, before, after in [
        ("Heading 1", 10.5, 5, 2),
        ("Heading 2", 9.8, 3, 1),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(31, 78, 121)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(18)
    title.font.bold = True
    title.font.color.rgb = RGBColor(31, 78, 121)
    title.paragraph_format.space_after = Pt(0)


def add_section_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.add_run(text.upper())
    p_pr = p._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F4E79")
    border.append(bottom)
    p_pr.append(border)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(0.8)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.08)
    p.add_run(text)


def add_two_col_row(table, left, right, bold_left=True):
    row = table.add_row().cells
    row[0].text = left
    row[1].text = right
    for i, cell in enumerate(row):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(cell, "FFFFFF")
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.font.size = Pt(9.2)
                if i == 0 and bold_left:
                    r.bold = True


def build():
    doc = Document()
    set_margins(doc)
    set_base_styles(doc)

    name = doc.add_paragraph(style="Title")
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.add_run("NIKKI SHUKLA")

    headline = doc.add_paragraph()
    headline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = headline.add_run("Fresher QA Tester | Manual Testing | Functional Testing | Bug Reporting")
    r.bold = True
    r.font.size = Pt(10.2)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.add_run("Mumbai, India | +91 8779132103 | nikkishukla95100@gmail.com | LinkedIn | GitHub")

    add_section_heading(doc, "Professional Summary")
    doc.add_paragraph(
        "Computer Engineering graduate (2026) seeking an entry-level QA Tester / Software Tester role. "
        "Good understanding of manual testing, SDLC, STLC, test case writing, defect reporting, and regression testing. "
        "Hands-on practice through an e-commerce testing project covering login, registration, search, cart, checkout, and order confirmation workflows."
    )

    add_section_heading(doc, "Skills")
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    add_two_col_row(table, "Testing", "Manual Testing, Functional Testing, Regression Testing, Smoke Testing, Negative Testing, UI Testing")
    add_two_col_row(table, "QA Documentation", "Test Scenarios, Test Cases, Test Plan, Bug Report, RTM, Test Execution Summary")
    add_two_col_row(table, "Defect Tracking", "JIRA basics / Mock JIRA, Severity, Priority, Bug Life Cycle")
    add_two_col_row(table, "Tools", "MS Excel, Google Sheets, Chrome DevTools basics, GitHub basics, VS Code")
    add_two_col_row(table, "Basic Technical", "HTML, CSS, JavaScript basics, SQL basics")

    add_section_heading(doc, "QA Project")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("E-Commerce Website Testing (Amazon/Flipkart Style)")
    run.bold = True
    p.add_run(" | Manual Testing Project")
    add_bullet(doc, "Created and executed 55 manual test cases for registration, login, product search, product details, add to cart, checkout, payment validation, and order confirmation.")
    add_bullet(doc, "Identified and documented 10 defects with bug ID, module, severity, priority, steps to reproduce, expected result, actual result, and status.")
    add_bullet(doc, "Performed functional, smoke, negative, UI, boundary value, and regression testing for important customer journeys.")
    add_bullet(doc, "Prepared test plan, test cases, bug report, mock JIRA tickets, requirement traceability matrix, and execution summary.")

    add_section_heading(doc, "Work Experience")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Full Stack Web Development Intern - Main Flow Services & Technologies Pvt. Ltd.")
    r.bold = True
    p.add_run(" | Jun 2025 - Dec 2025")
    add_bullet(doc, "Worked on basic web development tasks and gained understanding of web application flow, forms, authentication, and user interface behavior.")
    add_bullet(doc, "Tested forms and basic application flows during development, which helped build interest in QA and software testing.")

    add_section_heading(doc, "Education")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("B.E. Computer Engineering - Lokmanya Tilak College of Engineering, University of Mumbai")
    r.bold = True
    p.add_run(" | 2022 - 2026")
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r = p2.add_run("12th Science (PCM + CS) - IDEAL College of Arts, Science & Commerce, Mumbai")
    r.bold = True
    p2.add_run(" | 80% | 2020 - 2022")

    add_section_heading(doc, "Certifications")
    add_bullet(doc, "IBM SkillsBuild - Web Development Fundamentals")
    add_bullet(doc, "Udemy - Full Stack Web Development")
    add_bullet(doc, "GFG - AWS Cloud Practitioner Essentials")

    add_section_heading(doc, "Interview-ready Project Line")
    doc.add_paragraph(
        "I tested an e-commerce application as a fresher QA project, wrote 55 test cases, logged 10 defects, and prepared QA documents like test plan, bug report, RTM, and execution summary."
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
